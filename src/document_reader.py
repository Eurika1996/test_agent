import os


class DocumentReader:
    def __init__(self):
        pass

    def read(self, file_path):
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext in ['.md', '.markdown', '.txt']:
            return self._read_text(file_path)
        elif ext == '.docx':
            return self._read_docx(file_path)
        elif ext == '.pdf':
            return self._read_pdf(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {ext}（支持 .md/.txt/.docx/.pdf）")

    def _read_text(self, file_path):
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            with open(file_path, 'r', encoding='gbk') as f:
                return f.read()

    def _read_docx(self, file_path):
        try:
            from docx import Document
        except ImportError:
            raise ImportError("请安装 python-docx: pip install python-docx")

        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    paragraphs.append(' | '.join(cells))

        return '\n'.join(paragraphs)

    def _read_pdf(self, file_path):
        try:
            import pdfplumber
        except ImportError:
            raise ImportError("请安装 pdfplumber: pip install pdfplumber")

        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return '\n'.join(text_parts)


def chunk_text(text, max_chunk_size=4000):
    paragraphs = text.split('\n')
    chunks = []
    current = []
    current_len = 0

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        para_len = len(para)
        if current_len + para_len + 1 > max_chunk_size and current:
            chunks.append('\n'.join(current))
            current = [para]
            current_len = para_len
        else:
            current.append(para)
            current_len += para_len + 1

    if current:
        chunks.append('\n'.join(current))

    return chunks
